import streamlit as st
import pandas as pd
from reportlab.pdfgen import canvas
from reportlab.lib.pagesizes import A4
from reportlab.lib.units import cm
from docx import Document
from openpyxl import Workbook
from io import BytesIO
from config.i18n import t

def generar_pdf(titulo, contenido):
    buffer = BytesIO()
    p = canvas.Canvas(buffer, pagesize=A4)
    ancho, alto = A4
    p.setFont("Helvetica-Bold", 16)
    p.drawString(2*cm, alto - 2*cm, titulo)
    p.setFont("Helvetica", 11)
    y = alto - 3.5*cm
    for linea in contenido.split("\n"):
        if y < 2*cm:
            p.showPage()
            y = alto - 2*cm
        p.drawString(2*cm, y, linea)
        y -= 0.6*cm
    p.save()
    buffer.seek(0)
    return buffer

def generar_excel(datos_df):
    buffer = BytesIO()
    wb = Workbook()
    ws = wb.active
    ws.title = "Resultados"
    ws.append(list(datos_df.columns))
    for _, fila in datos_df.iterrows():
        ws.append(list(fila.values))
    wb.save(buffer)
    buffer.seek(0)
    return buffer

def generar_word(titulo, resumen, tabla_datos=None):
    doc = Document()
    doc.add_heading(titulo, 0)
    doc.add_paragraph(resumen)
    if tabla_datos is not None:
        doc.add_heading("Métricas", level=2)
        tabla = doc.add_table(rows=1, cols=len(tabla_datos.columns))
        for c, nombre in enumerate(tabla_datos.columns):
            tabla.rows[0].cells[c].text = str(nombre)
        for _, fila in tabla_datos.iterrows():
            fila_celdas = tabla.add_row().cells
            for c, valor in enumerate(fila):
                fila_celdas[c].text = str(valor)
    buffer = BytesIO()
    doc.save(buffer)
    buffer.seek(0)
    return buffer

def generar_reportes():
    st.header("📄 " + t("menu_reportes", st.session_state.idioma))
    tipo_reporte = st.selectbox("Formato de Salida", ["PDF", "Word", "Excel"])
    titulo = st.text_input("Título del Reporte", "Informe de Simulación Agronómica")
    resumen = st.text_area("Resumen / Observaciones", 
                           "Este informe presenta los resultados de la simulación de escenarios y evaluación de riesgo.")

    datos_demo = {
        "Zona": ["Z1", "Z2", "Z3", "Z4"],
        "Rendimiento Estimado (ton/ha)": [8.42, 7.15, 9.03, 6.88],
        "Riesgo de Pérdida (%)": [12.5, 41.3, 8.7, 72.1],
        "Intervención Sugerida": ["Riego suplementario", "Fertilización nitrogenada", 
                                   "Manejo estándar", "⚠️ Intervención URGENTE"]
    }
    df_demo = pd.DataFrame(datos_demo)
    
    st.subheader("📋 Datos del Reporte")
    st.dataframe(df_demo, use_container_width=True)

    if st.button("📥 Generar Reporte", type="primary"):
        nombre_base = titulo.replace(" ", "_")
        if tipo_reporte == "PDF":
            contenido = f"{titulo}\n\nFecha de generación: {pd.Timestamp.now().strftime('%d/%m/%Y %H:%M')}\n\n{resumen}\n\n"
            contenido += "=== Resumen de Resultados ===\n"
            contenido += df_demo.to_string(index=False)
            archivo = generar_pdf(titulo, contenido)
            st.download_button("⬇️ Descargar PDF", archivo, file_name=f"{nombre_base}.pdf", mime="application/pdf")
        elif tipo_reporte == "Excel":
            archivo = generar_excel(df_demo)
            st.download_button("⬇️ Descargar Excel", archivo, file_name=f"{nombre_base}.xlsx",
                               mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
        elif tipo_reporte == "Word":
            archivo = generar_word(titulo, resumen, df_demo)
            st.download_button("⬇️ Descargar Word", archivo, file_name=f"{nombre_base}.docx",
                               mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document")
